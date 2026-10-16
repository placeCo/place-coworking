#!/usr/bin/env python3
"""Fill Place Co. payment vouchers (ใบสำคัญจ่าย) from CSV using the template.

Usage:
  python3 fill_voucher.py vouchers.csv
  python3 fill_voucher.py vouchers.csv --out-dir ./out
  python3 fill_voucher.py vouchers.csv --no-round          # keep satang
  python3 fill_voucher.py --demo                          # write one example

CSV columns (header row required):
  date, payee, method, desc, amount, inv_no
Optional: voucher_no, remarks, cheque_no, bank, note, amount_words
WHT (optional, per voucher, read from the first line): wht_type, vat, amount_ex_vat, payee_type, wht_bearer
  wht_type: rent | service | hire | advertising | transport | none (empty = none, e.g. goods)
  base = pre-VAT amount: amount_ex_vat if given, else (sum of lines - vat) if vat given, else sum of lines.
  wht_bearer: payee (default: withheld from the payment) | company (gross-up: company bears the tax)
  payee_type: individual | company (default guessed from the payee name) -> PND3 / PND53
Several CSV rows with the same voucher_no = ONE voucher with several lines
(one payment / one slip -> one voucher; total = sum of lines, max 8 lines).

method: Cash | Transfer | Cheque  (also Thai: เงินสด | เงินโอน | เช็ค)
amount: number; by default rounded to whole baht (office practice).
Signatures are NEVER filled — leave for wet ink after print.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "payment_voucher_template.xlsx"

ONES = ["", "หนึ่ง", "สอง", "สาม", "สี่", "ห้า", "หก", "เจ็ด", "แปด", "เก้า"]
UNITS = ["", "สิบ", "ร้อย", "พัน", "หมื่น", "แสน", "ล้าน"]

# ---- Withholding tax (ภาษีหัก ณ ที่จ่าย) — rules confirmed by Pat (accountant), LINE 10.10.2026 17:12, point 4 ----
# 4.1 rent: 5% of the amount BEFORE VAT (if any) · 4.2 hire/services: 3% when the payment is 1,000 THB or more
# 4.3 advertising: 2% · 4.4 transport: 1%
# Each withholding -> 50 ทวิ certificate in 2 copies: copy 1 attached to the payment voucher, copy 2 kept (สำเนา).
# Pat files PND3 (individuals) / PND53 (companies) and draws the remittance. Goods purchases: no WHT.
WHT_RULES = {
    #  type          rate                min base (pre-VAT)   Thai label
    "rent":        (Decimal("0.05"), Decimal("0"),    "ค่าเช่า"),
    "service":     (Decimal("0.03"), Decimal("1000"), "ค่าบริการ"),
    "hire":        (Decimal("0.03"), Decimal("1000"), "ค่าจ้าง"),
    "advertising": (Decimal("0.02"), Decimal("0"),    "ค่าโฆษณา"),
    "transport":   (Decimal("0.01"), Decimal("0"),    "ค่าขนส่ง"),
}
WHT_ALIASES = {"ค่าเช่า": "rent", "บริการ": "service", "ค่าบริการ": "service", "ค่าจ้าง": "hire", "จ้าง": "hire",
               "โฆษณา": "advertising", "ค่าโฆษณา": "advertising", "ขนส่ง": "transport", "ค่าขนส่ง": "transport",
               "services": "service", "ads": "advertising", "": "none", "none": "none", "goods": "none", "-": "none"}
NOTE_50TAWI = ("ออกหนังสือรับรองการหักภาษี ณ ที่จ่าย (50 ทวิ) 2 ฉบับ: ฉบับที่ 1 แนบใบสำคัญจ่ายนี้ · "
               "ฉบับที่ 2 เก็บเป็นสำเนา · ผู้ทำบัญชียื่น {form}")


def _d(x) -> Decimal | None:
    x = (str(x).replace(",", "").strip() if x is not None else "")
    return Decimal(x) if x else None


def compute_wht(head: dict, gross_total: Decimal) -> dict | None:
    """WHT for one voucher. gross_total = sum of the voucher lines (what we owe incl. VAT).
    Returns None when no WHT applies, else dict(type, rate, base, wht, net, bearer, form, label)."""
    t = (head.get("wht_type") or "").strip().lower()
    t = WHT_ALIASES.get(t, t)
    if t == "none":
        return None
    if t not in WHT_RULES:
        raise SystemExit(f"voucher {head.get('voucher_no')}: unknown wht_type {t!r} (use {sorted(WHT_RULES)} or none)")
    rate, min_base, label = WHT_RULES[t]
    base = _d(head.get("amount_ex_vat"))
    if base is None:
        vat = _d(head.get("vat"))
        base = gross_total - vat if vat is not None else gross_total
    bearer = (head.get("wht_bearer") or "payee").strip().lower()
    ptype = (head.get("payee_type") or "").strip().lower()
    if not ptype:
        nm = head.get("payee") or ""
        ptype = "company" if any(k in nm for k in ("บริษัท", "ห้างหุ้นส่วน", "หจก", "Co.", "CO.", "Ltd", "LTD")) else "individual"
    form = "ภ.ง.ด.53" if ptype == "company" else "ภ.ง.ด.3"
    if bearer == "company":
        # gross-up: payee receives the full agreed amount; company bears the tax -> base = agreed / (1 - rate)
        base = (base / (1 - rate)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if base < min_base:
        return None
    wht = (base * rate).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    net = gross_total if bearer == "company" else gross_total - wht
    return dict(type=t, rate=rate, base=base, wht=wht, net=net, bearer=bearer, form=form, label=label, payee_type=ptype)


def _read_group(n: int) -> str:
    if n == 0:
        return ""
    text = str(n)
    s = ""
    L = len(text)
    for i, ch in enumerate(text):
        d = int(ch)
        pos = L - i - 1
        if d == 0:
            continue
        if pos == 0:
            if d == 1 and L > 1 and text[-2] != "0":
                s += "เอ็ด"
            else:
                s += ONES[d]
        elif pos == 1:
            if d == 1:
                s += "สิบ"
            elif d == 2:
                s += "ยี่สิบ"
            else:
                s += ONES[d] + "สิบ"
        else:
            s += ONES[d] + UNITS[pos]
    return s


def baht_text(amount) -> str:
    a = Decimal(str(amount)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    neg = a < 0
    a = abs(a)
    baht = int(a)
    satang = int((a - baht) * 100)

    def chunk(n: int) -> str:
        if n == 0:
            return "ศูนย์"
        parts = []
        million = n // 1_000_000
        rest = n % 1_000_000
        if million:
            parts.append(_read_group(million) + "ล้าน")
        if rest:
            parts.append(_read_group(rest))
        elif not million:
            parts.append("ศูนย์")
        return "".join(parts)

    words = chunk(baht) + "บาท"
    words += "ถ้วน" if satang == 0 else _read_group(satang) + "สตางค์"
    return ("ลบ" + words) if neg else words


def round_baht(amount, do_round: bool = True) -> Decimal:
    a = Decimal(str(amount))
    if do_round:
        return a.quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    return a.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _chk(method: str, kind: str) -> bool:
    m = (method or "").strip().lower()
    aliases = {
        "cash": {"cash", "เงินสด", "c", "เงินสด (ชำระปลายทาง)"},
        "transfer": {"transfer", "เงินโอน", "โอน", "bank", "bank transfer", "t"},
        "cheque": {"cheque", "check", "เช็ค", "เช็คธนาคาร", "q"},
    }
    return m in aliases[kind]


def fill_sheet(ws, row, do_round: bool = True) -> None:
    """Fill one voucher. `row` is a dict (one line) or a list of dicts (several lines,
    same voucher_no: header fields taken from the first line; up to 8 lines)."""
    lines = row if isinstance(row, list) else [row]
    if len(lines) > 8:
        raise SystemExit(f"voucher {lines[0].get('voucher_no')}: max 8 lines, got {len(lines)}")
    head = lines[0]
    fmt = "#,##0.00" if not do_round else "#,##0"
    method = head.get("method") or ""
    ws["F3"] = head.get("voucher_no") or ""
    ws["F4"] = head.get("date") or ""
    ws["B6"] = head.get("payee") or ""
    # checkbox labels come from the template (bilingual or Thai-only template)
    for cell, kind, dflt in (("B7", "cash", "เงินสด / Cash"), ("D7", "transfer", "เงินโอน / Transfer"), ("F7", "cheque", "เช็คธนาคาร / Cheque")):
        lbl = str(ws[cell].value or "").lstrip("☐☑ ").strip() or dflt
        ws[cell] = f"{'☑' if _chk(method, kind) else '☐'} {lbl}"
    ws["B8"] = head.get("remarks") or ""
    ws["E8"] = head.get("cheque_no") or ""
    ws["G8"] = head.get("bank") or ""

    # clear description lines 11..18
    for i in range(8):
        r = 11 + i
        ws[f"A{r}"] = ""
        ws[f"B{r}"] = ""
        ws[f"F{r}"] = None
        ws[f"H{r}"] = ""

    total = Decimal("0")
    for i, ln in enumerate(lines):
        r = 11 + i
        amt = round_baht(ln["amount"], do_round)
        total += amt
        ws[f"A{r}"] = ln.get("inv_no") or ""
        ws[f"B{r}"] = ln.get("desc") or ""
        ws[f"F{r}"] = float(amt)
        ws[f"F{r}"].number_format = fmt
        ws[f"H{r}"] = ln.get("note") or ""
        if i > 0:  # copy row-11 styling (wrap/fonts) to extra lines
            for c in "ABFH":
                ws[f"{c}{r}"].font = copy(ws[f"{c}11"].font)
                ws[f"{c}{r}"].alignment = copy(ws[f"{c}11"].alignment)
            ws.row_dimensions[r].height = ws.row_dimensions[11].height
    if len(lines) > 1:
        for i in range(len(lines)):
            ws.row_dimensions[11 + i].height = 40

    ws["F19"] = float(total)
    ws["F19"].number_format = fmt
    w = compute_wht(head, total)
    pay = total
    if w:
        pct = f"{(w['rate'] * 100).normalize()}%"
        ws.merge_cells("A20:E20") if "A20:E20" not in {str(m) for m in ws.merged_cells.ranges} else None
        if w["bearer"] == "company":
            ws["A20"] = (f"ภาษีหัก ณ ที่จ่าย {w['label']} {pct} บริษัทออกภาษีแทน (ฐาน {w['base']:,.2f} บาท) "
                         f"— ผู้รับเงินได้รับเต็มจำนวน")
            ws["F20"] = float(w["wht"])
            ws["H20"] = "ภาระของบริษัท"
        else:
            ws["A20"] = f"หัก ภาษีหัก ณ ที่จ่าย {w['label']} {pct} (ฐานก่อน VAT {w['base']:,.2f} บาท) · จ่ายสุทธิ {w['net']:,.2f} บาท"
            ws["F20"] = float(-w["wht"])
            ws["H20"] = f"จ่ายสุทธิ {w['net']:,.2f}"
            pay = w["net"]
        ws["F20"].number_format = "#,##0.00;-#,##0.00"
        for c in ("A20", "F20", "H20"):
            ws[c].font = copy(ws["A19"].font); ws[c].alignment = copy(ws["B11"].alignment)
        ws["F20"].alignment = copy(ws["F19"].alignment)
        ws.row_dimensions[20].height = 30
        ws.merge_cells("A24:H24") if "A24:H24" not in {str(m) for m in ws.merged_cells.ranges} else None
        ws["A24"] = NOTE_50TAWI.format(form=w["form"])
        ws["A24"].font = copy(ws["A23"].font); ws["A24"].alignment = copy(ws["B11"].alignment)
    ws["A22"] = head.get("amount_words") or baht_text(pay)
    # signatures intentionally untouched


def safe_name(s: str) -> str:
    s = re.sub(r"[^\w\-]+", "_", (s or "voucher").strip(), flags=re.UNICODE)
    return (s[:40] or "voucher").strip("_")


def process_csv(csv_path: Path, out_dir: Path, do_round: bool, template: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    with csv_path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        required = {"date", "payee", "method", "desc", "amount", "inv_no"}
        if not reader.fieldnames or not required.issubset({h.strip() for h in reader.fieldnames}):
            raise SystemExit(
                f"CSV needs columns: {sorted(required)}. Got: {reader.fieldnames}"
            )
        groups: dict = {}
        order = []
        wht_rows: list = []
        for idx, raw in enumerate(reader, start=1):
            row = {(k or "").strip(): (v.strip() if isinstance(v, str) else v) for k, v in raw.items()}
            if not row.get("payee") and not row.get("amount"):
                continue
            key = row.get("voucher_no") or f"{idx:03d}"
            row.setdefault("voucher_no", key)
            if key not in groups:
                groups[key] = []
                order.append(key)
            groups[key].append(row)  # rows sharing voucher_no = one voucher, several lines
        for vn in order:
            lines = groups[vn]
            head = lines[0]
            wb = load_workbook(template)
            ws = wb["Payment Voucher"]
            fill_sheet(ws, lines, do_round=do_round)
            w = compute_wht(head, sum(round_baht(l["amount"], do_round) for l in lines))
            if w:
                wht_rows.append({"voucher_no": vn, "date": head.get("date"), "payee": head.get("payee"),
                                 "payee_type": w["payee_type"], "wht_type": w["type"], "rate": str(w["rate"]),
                                 "base_ex_vat": f"{w['base']:.2f}", "wht": f"{w['wht']:.2f}", "net_paid": f"{w['net']:.2f}",
                                 "bearer": w["bearer"], "form": w["form"], "cert_50tawi": "2 copies: 1 attached to voucher, 1 kept"})
            fname = f"PV-{safe_name(str(head.get('date')))}-{safe_name(str(head.get('payee')))}-{vn}.xlsx"
            dest = out_dir / fname
            wb.save(dest)
            written.append(dest)
            print(f"Wrote {dest}" + (f" ({len(lines)} lines)" if len(lines) > 1 else ""))
    if wht_rows:
        reg = out_dir / "WHT-50TAWI-REGISTER.csv"
        with reg.open("w", newline="", encoding="utf-8") as f:
            wr = csv.DictWriter(f, fieldnames=list(wht_rows[0])); wr.writeheader(); wr.writerows(wht_rows)
        print(f"Wrote {reg} ({len(wht_rows)} withholding(s) -> 50 ทวิ x2 each)")
    return written


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Fill Place Co. payment vouchers from CSV")
    ap.add_argument("csv", nargs="?", help="Input CSV path")
    ap.add_argument("--out-dir", default=str(HERE / "out"), help="Output folder")
    ap.add_argument("--no-round", action="store_true", help="Keep satang (default: round to whole baht)")
    ap.add_argument("--demo", action="store_true", help="Write PP Produce demo row")
    ap.add_argument("--template", default=str(TEMPLATE), help="Template xlsx path")
    args = ap.parse_args(argv)

    template = Path(args.template)
    if not template.exists():
        print(f"Template not found: {template}", file=sys.stderr)
        return 1

    if args.demo:
        demo = HERE / "vouchers.example.csv"
        demo.write_text(
            "date,payee,method,desc,amount,inv_no,note\n"
            "23 Sep 2026,บริษัท พี.พี. ผลิตผล จำกัด / PP Produce,Transfer,"
            "INGREDIENTS FOR KITCHEN,1454.74,B210411175,rounded on voucher\n",
            encoding="utf-8",
        )
        print(f"Wrote {demo}")
        process_csv(demo, Path(args.out_dir), do_round=not args.no_round, template=template)
        return 0

    if not args.csv:
        ap.print_help()
        return 2
    process_csv(Path(args.csv), Path(args.out_dir), do_round=not args.no_round, template=template)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
